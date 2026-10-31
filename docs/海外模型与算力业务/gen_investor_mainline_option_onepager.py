#!/usr/bin/env python3
"""Investor one-pager: mainline (Baidu compute/token) + option (AI service SKUs)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("投资人一页纸_主线加期权.pptx")
OUT_MD = Path(__file__).with_name("投资人一页纸_主线加期权.md")

INK = RGBColor(0x0F, 0x1C, 0x2E)
INK2 = RGBColor(0x1A, 0x2F, 0x45)
MUTED = RGBColor(0x6B, 0x7C, 0x8E)
CARD = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COPPER = RGBColor(0xC4, 0x5C, 0x26)
TEAL = RGBColor(0x1F, 0x7A, 0x6E)
AMBER = RGBColor(0xB8, 0x7A, 0x1E)


def set_run(run, text, size=14, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"


def add_text(tf, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, space_after=4):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)


def clear_tf(tf):
    tf.clear()
    tf.word_wrap = True


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def band(slide, top, height, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), top, Inches(13.333), height)
    fill_shape(sh, color)


def card(slide, left, top, width, height, fill=CARD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(sh, fill)
    sh.adjustments[0] = 0.06
    return sh


def textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def write_md():
    OUT_MD.write_text(
        """# 投资人一页纸：主线 + 期权

**主体**：香港公司 · **供给**：国产主流模型官方折扣 + 百度智能云 · **阶段**：Pre-revenue

---

## 一句话

用**官方折扣成本**与**百度供给**，经香港公司向海外卖 AI 能力；主线做**算力 / Token 渠道**，期权做**可复制的 AI 服务标品**——两门生意都留，但资源与叙事分清。

---

## 结构：主线 + 期权（不是二选一）

| | 主线（现在投精力） | 期权（有客户名单再放大） |
|---|---|---|
| 卖什么 | 模型调用额度 / 百度云算力 | 可配置的业务结果（标品订阅） |
| 客户 | 海外开发者、技术型小团队 → 再企业 | 已有行业/外贸客户（信任已在） |
| 我们出什么 | 港司收款、渠道、一线支持 | 场景理解、本地化交付、续费 |
| 百度出什么 | 模型/算力、账户与渠道政策 | 智能体工具与模型支撑 |
| 为什么是主线 | 已有折扣与百度关系，轻资产可启动 | 毛利厚，但取决于是否真有可复用场景 |
| 3–6 个月见钱 | 取决于书面授权与结算能否谈成 | 样板客户手工交付即可收费 |
| 主要风险 | 被官方聚合平台压价、授权谈不拢 | 滑回外包、做不成标品 |

**资源配比建议（90 天）**：主线 **7** · 期权 **3**（先确认百度政策与报价；期权只做一个「标品一号」样板，不铺第二条产品线）。

---

## 主线：算力 / Token 渠道（渠道生意，不是灰 Key）

1. 境外主体开户与结算（港司、币种、发票）  
2. 区域渠道 / 分销折扣是否允许跨境转售  
3. 子账号落户（优先）vs 纯返点  
4. 跨境数据与内容合规书面口径  
5. 延迟 / 节点 / SLA  

**定价锚**：售价 ≈ 结算成本 × **1.5～1.8** · 预充值 · 忌裸卖共享 Key。

**备用**：百度路径谈不拢时，启用 One API 自建网关（非当前主投入）。

---

## 期权：AI 服务标品（渠道生意的加厚）

验收标准只有一条：第 10 个客户上线是否「开账号 + 配置」，而不是再开一轮开发。

候选（有真实客户抱怨再选 1 个）：多语种图册/详情页 · 询盘邮件与报价校对 · 单证合同条款比对 · 客户检索与会前简报。

路径：样板手工跑通 → 抽 80% 共性固化 → 复制到第 3/5/10 家。可对技术型客户**附带**卖调用额度，额度是增值项不是主业叙事。

---

## 投资人看什么

| 核验点 | 通过标准 |
|---|---|
| 百度书面授权 | 身份、结算、落户、出海、品牌五句有记录 |
| 单位经济 | 含支付与客服后仍有可接受贡献利润 |
| 主线种子单 | 3–5 个可对账用量的付费/预付账户 |
| 期权真实性 | 「标品一号」来自真实重复痛点，而非空想清单 |
| 红线 | 无灰通道、无无限包月、无未授权跨境转售 |

---

## 结论

这是一桩 **「授权渠道」** 生意：主线用折扣与百度供给换海外经常性调用收入；期权用客户信任把 AI 做成可订阅结果。投的是**分销与场景能力**，不是自建机房。下一道闸门是**百度书面条款**，不是再多买服务器。

*说明：财务数字若引用规划情景，均为假设非实账；供应侧政策以书面确认为准。*
""",
        encoding="utf-8",
    )


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band(slide, Inches(0), Inches(7.5), WHITE)
    band(slide, Inches(0), Inches(0.08), COPPER)
    band(slide, Inches(0), Inches(0.95), INK)

    tb = textbox(slide, Inches(0.55), Inches(0.22), Inches(10), Inches(0.55))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(tf, "投资人一页纸 · 主线 + 期权", size=22, bold=True, color=WHITE, space_after=0)

    tb = textbox(slide, Inches(0.55), Inches(1.1), Inches(12.2), Inches(0.55))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(
        tf,
        "官方折扣 × 香港收款 × 百度供给 → 海外卖 AI 能力。主线做算力/Token 渠道；期权做可复制服务标品。两门都留，资源分清。",
        size=13,
        color=INK2,
        space_after=0,
    )

    # Mainline card
    card(slide, Inches(0.45), Inches(1.75), Inches(6.15), Inches(3.55))
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(1.75), Inches(6.15), Inches(0.42))
    fill_shape(bar, TEAL)
    tb = textbox(slide, Inches(0.65), Inches(1.82), Inches(5.7), Inches(0.32))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(tf, "主线  ·  算力 / Token 渠道（现在）", size=14, bold=True, color=WHITE, space_after=0)

    lines_main = [
        "卖：模型调用额度 / 百度云算力（按量）",
        "客：海外开发者与小团队 → 再企业",
        "我们：港司收款、一线支持、渠道",
        "百度：模型/算力、账户与渠道政策",
        "定价：成本 × 1.5～1.8 · 预充值 · 忌灰 Key",
        "闸门：书面授权（开户/折扣/落户/出海/品牌）",
        "备用：谈不拢再启 One API，非当前主投入",
    ]
    tb = textbox(slide, Inches(0.65), Inches(2.3), Inches(5.7), Inches(2.8))
    tf = tb.text_frame
    clear_tf(tf)
    for i, line in enumerate(lines_main):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        run = p.add_run()
        set_run(run, "•  " + line, size=12, color=INK2)

    # Option card
    card(slide, Inches(6.75), Inches(1.75), Inches(6.15), Inches(3.55))
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.75), Inches(1.75), Inches(6.15), Inches(0.42))
    fill_shape(bar, AMBER)
    tb = textbox(slide, Inches(6.95), Inches(1.82), Inches(5.7), Inches(0.32))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(tf, "期权  ·  AI 服务标品（有客户再放大）", size=14, bold=True, color=WHITE, space_after=0)

    lines_opt = [
        "卖：业务结果订阅（开账号+配置即用）",
        "客：已有行业/外贸客户（信任已在）",
        "验收：第10家仍只需配置，不做外包",
        "路径：样板手工 → 固化80% → 复制放大",
        "候选：图册/询盘/单证/会前简报（四选一）",
        "与主线关系：可附带卖额度，额度是增值项",
        "见钱：样板客户手工交付即可收费",
    ]
    tb = textbox(slide, Inches(6.95), Inches(2.3), Inches(5.7), Inches(2.8))
    tf = tb.text_frame
    clear_tf(tf)
    for i, line in enumerate(lines_opt):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        run = p.add_run()
        set_run(run, "•  " + line, size=12, color=INK2)

    # Bottom strip: allocation + diligence
    card(slide, Inches(0.45), Inches(5.45), Inches(4.0), Inches(1.7))
    tb = textbox(slide, Inches(0.6), Inches(5.55), Inches(3.7), Inches(1.45))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(tf, "90 天资源", size=12, bold=True, color=TEAL, space_after=6)
    add_text(tf, "主线 7  ·  期权 3", size=18, bold=True, color=INK, space_after=6)
    add_text(tf, "期权只打磨 1 个标品一号", size=11, color=MUTED, space_after=0)

    card(slide, Inches(4.6), Inches(5.45), Inches(8.3), Inches(1.7))
    tb = textbox(slide, Inches(4.8), Inches(5.55), Inches(7.9), Inches(1.45))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(tf, "投资人核验（通过才加码）", size=12, bold=True, color=COPPER, space_after=6)
    add_text(
        tf,
        "①百度书面授权  ②含支付/客服后仍有贡献利润  ③主线 3–5 个可对账种子单  ④期权痛点来自真实客户  ⑤无灰通道",
        size=12,
        color=INK2,
        space_after=0,
    )

    prs.save(OUT)
    print(f"Wrote {OUT}")


def main():
    write_md()
    build_ppt()
    print(f"Wrote {OUT_MD}")


if __name__ == "__main__":
    main()
