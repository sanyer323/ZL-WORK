#!/usr/bin/env python3
"""Investor-facing PPT export of the overseas AI compute dashboard."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("投资者看板_海外AI算力渠道.pptx")

INK = RGBColor(0x0F, 0x1C, 0x2E)
INK2 = RGBColor(0x1A, 0x2F, 0x45)
MUTED = RGBColor(0x6B, 0x7C, 0x8E)
CARD = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COPPER = RGBColor(0xC4, 0x5C, 0x26)
TEAL = RGBColor(0x1F, 0x7A, 0x6E)
BAD = RGBColor(0xA8, 0x3B, 0x2C)


def set_run(run, text, size=18, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"


def add_text(tf, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, space_after=6):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)


def clear_tf(tf):
    tf.clear()
    tf.word_wrap = True


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def band(slide, top, height, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), top, Inches(13.333), height)
    fill_shape(sh, color)
    return sh


def accent_bar(slide, left, top, width=Inches(0.12), height=Inches(0.55)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(sh, COPPER)
    return sh


def card(slide, left, top, width, height, fill=CARD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(sh, fill)
    sh.adjustments[0] = 0.08
    return sh


def textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def section_header(slide, title, subtitle=None):
    band(slide, Inches(0), Inches(0.08), COPPER)
    accent_bar(slide, Inches(0.7), Inches(0.35))
    tb = textbox(slide, Inches(1.0), Inches(0.32), Inches(11.5), Inches(0.6))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(tf, title, size=26, bold=True, color=INK, space_after=0)
    if subtitle:
        tb2 = textbox(slide, Inches(1.0), Inches(0.95), Inches(11.5), Inches(0.4))
        tf2 = tb2.text_frame
        clear_tf(tf2)
        add_text(tf2, subtitle, size=13, color=MUTED, space_after=0)


def bullet_block(slide, left, top, width, height, lines, size=15):
    tb = textbox(slide, left, top, width, height)
    tf = tb.text_frame
    clear_tf(tf)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        set_run(run, line, size=size, color=INK2)


def table_slide(slide, headers, rows, left=Inches(0.7), top=Inches(1.55), width=Inches(11.9), row_h=0.55):
    cols = len(headers)
    nrows = len(rows) + 1
    table_shape = slide.shapes.add_table(nrows, cols, left, top, width, Inches(row_h * nrows))
    table = table_shape.table
    col_w = int(width / cols)
    for i in range(cols):
        table.columns[i].width = col_w
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                set_run(r, h, size=12, bold=True, color=WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    set_run(r, val, size=12, color=INK2)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else CARD


def title_slide(prs, title, subtitle, foot):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band(slide, Inches(0), Inches(7.5), INK)
    band(slide, Inches(0), Inches(0.18), COPPER)
    tb = textbox(slide, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.4))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(tf, title, size=34, bold=True, color=WHITE, space_after=10)
    tb2 = textbox(slide, Inches(0.9), Inches(3.7), Inches(11.5), Inches(1.2))
    tf2 = tb2.text_frame
    clear_tf(tf2)
    add_text(tf2, subtitle, size=18, color=RGBColor(0xB8, 0xC7, 0xD6), space_after=8)
    tb3 = textbox(slide, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.5))
    tf3 = tb3.text_frame
    clear_tf(tf3)
    add_text(tf3, foot, size=12, color=MUTED)


def two_col_cards(slide, left_title, left_lines, right_title, right_lines, left_color=TEAL, right_color=COPPER):
    card(slide, Inches(0.7), Inches(1.6), Inches(5.7), Inches(5.0))
    card(slide, Inches(6.9), Inches(1.6), Inches(5.7), Inches(5.0))
    s1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.6), Inches(5.7), Inches(0.55))
    fill_shape(s1, left_color)
    s2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.7), Inches(0.55))
    fill_shape(s2, right_color)
    t1 = textbox(slide, Inches(0.95), Inches(1.68), Inches(5.2), Inches(0.4))
    tf = t1.text_frame
    clear_tf(tf)
    add_text(tf, left_title, size=16, bold=True, color=WHITE, space_after=0)
    t2 = textbox(slide, Inches(7.15), Inches(1.68), Inches(5.2), Inches(0.4))
    tf = t2.text_frame
    clear_tf(tf)
    add_text(tf, right_title, size=16, bold=True, color=WHITE, space_after=0)
    bullet_block(slide, Inches(0.95), Inches(2.4), Inches(5.2), Inches(3.9), left_lines, size=14)
    bullet_block(slide, Inches(7.15), Inches(2.4), Inches(5.2), Inches(3.9), right_lines, size=14)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_slide(
        prs,
        "海外 AI 模型 / 算力渠道",
        "投资者看板 · Investor Brief",
        "香港公司主体 · 官方折扣供给 · 百度承载交付 · Pre-revenue 规划情景",
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "数据口径", "本材料财务数字为规划情景，不是历史实账")
    bullet_block(
        s,
        Inches(1.0),
        Inches(1.8),
        Inches(11),
        Inches(4.5),
        [
            "• 倍率、客户数、ARPU 为内部假设，用于敏感度讨论",
            "• 落地后以百度结算单与 Stripe / 香港公司实收替换",
            "• 当前阶段重点：书面授权与供给，而非自建机房",
            "• 明确不做：共享账号、无限包月、无授权跨境灌量",
        ],
        size=18,
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "一句话投资论点")
    card(s, Inches(0.7), Inches(1.7), Inches(11.9), Inches(2.2))
    tb = textbox(s, Inches(1.0), Inches(2.1), Inches(11.3), Inches(1.5))
    tf = tb.text_frame
    clear_tf(tf)
    add_text(
        tf,
        "官方成本优势 × 香港收款 × 百度供给 = 跨境 AI 渠道生意。",
        size=22,
        bold=True,
        color=INK,
        align=PP_ALIGN.CENTER,
        space_after=10,
    )
    add_text(
        tf,
        "投的是授权与分销能力，不是自建算力资产。",
        size=18,
        color=COPPER,
        align=PP_ALIGN.CENTER,
    )
    cards = [
        ("供给", "官方折扣\n主流国产模型可覆盖"),
        ("主体", "香港公司\nUSD 收款与签约"),
        ("交付", "百度承载\n不买服务器"),
        ("阶段", "Pre-revenue\n钉死授权再放量"),
    ]
    for i, (t, body) in enumerate(cards):
        left = Inches(0.7 + i * 3.1)
        card(s, left, Inches(4.3), Inches(2.9), Inches(2.4))
        tb = textbox(s, left + Inches(0.2), Inches(4.5), Inches(2.5), Inches(0.4))
        tf = tb.text_frame
        clear_tf(tf)
        add_text(tf, t, size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER, space_after=0)
        tb2 = textbox(s, left + Inches(0.2), Inches(5.1), Inches(2.5), Inches(1.3))
        tf2 = tb2.text_frame
        clear_tf(tf2)
        for j, line in enumerate(body.split("\n")):
            p = tf2.paragraphs[0] if j == 0 else tf2.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.space_after = Pt(4)
            run = p.add_run()
            set_run(run, line, size=13, color=INK2)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "投资者该看什么", "四根杠杆")
    table_slide(
        s,
        ["杠杆", "我们有什么", "投资者核验点"],
        [
            ["成本", "官方折扣（非灰色 Key）", "合同是否允许境外转售"],
            ["信任", "香港公司 + 可挂 Baidu Cloud", "书面渠道/代理身份"],
            ["运营杠杆", "轻资产，百度出资源", "子账号转售能否谈成"],
            ["扩张", "模型 API → GPU 可同渠道延伸", "首年是否可复核用量"],
        ],
        top=Inches(1.8),
        row_h=0.9,
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "商业模式与资金流", "轻资产渠道，而非买卡转租")
    table_slide(
        s,
        ["环节", "谁", "动作"],
        [
            ["获客/签约", "我方（香港公司）", "海外开发者与小团队"],
            ["收款", "香港公司", "USD / Stripe 等"],
            ["进货结算", "我方 → 百度/厂商", "官方折扣价"],
            ["交付", "百度云 / 模型 API", "GPU 或推理配额"],
            ["一线支持", "我方", "账号、账单、续费"],
            ["底层故障", "百度", "平台与资源可用性"],
        ],
        top=Inches(1.6),
        row_h=0.7,
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "规划情景 · 12 个月（示意）", "售价 ≈ 结算成本 × 1.5～1.8")
    table_slide(
        s,
        ["情景", "倍率", "M12 客户", "ARPU/月", "示意毛利率"],
        [
            ["保守", "1.50×", "80", "$420", "33%"],
            ["基准", "1.65×", "140", "$680", "39%"],
            ["进取", "1.80×", "220", "$950", "44%"],
        ],
        top=Inches(1.6),
        row_h=0.85,
    )
    note = textbox(s, Inches(0.7), Inches(5.0), Inches(11.9), Inches(1.8))
    tf = note.text_frame
    clear_tf(tf)
    add_text(tf, "基准情景粗算（M12 月）：", size=15, bold=True, color=INK, space_after=8)
    add_text(tf, "收入 ≈ 140 × $680 ≈ $95,200 / 月", size=14, color=INK2, space_after=4)
    add_text(tf, "毛利 ≈ 收入 × 39% ≈ $37,500 / 月（未扣获客与合规摊销）", size=14, color=INK2, space_after=4)
    add_text(tf, "贡献利润另扣支付约 3.5% + 客服约 6%", size=14, color=MUTED, space_after=0)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "单位经济（成本=100 归一）", "以基准 1.65× 为例")
    table_slide(
        s,
        ["项目", "金额（USD）", "说明"],
        [
            ["结算成本", "100", "百度/厂商官方折扣进价"],
            ["售价", "165", "成本 × 1.65"],
            ["毛利空间", "65", "毛利率约 39%"],
            ["支付摩擦", "~6", "约售价 3.5%"],
            ["客服余量", "~10", "约售价 6%"],
            ["贡献利润", "~49", "未含获客与合规摊销"],
        ],
        top=Inches(1.55),
        row_h=0.7,
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "收入结构目标（示意）", "首年以推理 API 获客")
    cards = [
        ("55%", "模型推理 API", "开发者入口、OpenAI 兼容体验"),
        ("30%", "GPU / 包月算力", "随百度授权加深放大"),
        ("15%", "企业配额与支持", "后期高客单"),
    ]
    for i, (n, t, d) in enumerate(cards):
        left = Inches(0.7 + i * 4.1)
        card(s, left, Inches(1.8), Inches(3.9), Inches(4.5))
        tb = textbox(s, left + Inches(0.25), Inches(2.3), Inches(3.4), Inches(0.8))
        tf = tb.text_frame
        clear_tf(tf)
        add_text(tf, n, size=36, bold=True, color=TEAL if i == 0 else COPPER if i == 1 else INK, align=PP_ALIGN.CENTER, space_after=0)
        tb2 = textbox(s, left + Inches(0.25), Inches(3.4), Inches(3.4), Inches(0.5))
        tf2 = tb2.text_frame
        clear_tf(tf2)
        add_text(tf2, t, size=16, bold=True, color=INK, align=PP_ALIGN.CENTER, space_after=0)
        tb3 = textbox(s, left + Inches(0.25), Inches(4.2), Inches(3.4), Inches(1.5))
        tf3 = tb3.text_frame
        clear_tf(tf3)
        add_text(tf3, d, size=14, color=MUTED, align=PP_ALIGN.CENTER)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "关键路径", "下一道闸门是百度书面条款")
    table_slide(
        s,
        ["阶段", "目标", "完成标准"],
        [
            ["W0–1", "百度五句话落地", "身份/结算/落户/出海/品牌有书面记录"],
            ["W2–3", "可卖清单 + 报价骨架", "GPU 与模型折扣表、进售价两列"],
            ["W4–6", "种子客户试单", "3–5 个付费或预付测试账户"],
            ["Q2", "稳定月结与复购", "可复核用量账单 + 毛利追踪"],
            ["备用", "One API 自建网关", "仅当百度路径谈不拢时启用"],
        ],
        top=Inches(1.6),
        row_h=0.75,
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "风险与缓释")
    table_slide(
        s,
        ["风险", "影响", "缓释"],
        [
            ["上游禁止境外转售", "模式失效", "先拿书面授权再放量"],
            ["配额/账号波动", "客户流失", "多模型备份 + 信用补偿"],
            ["支付拒付与盗刷", "毛利被吞", "预充值、限额、人工开户"],
            ["纯差价无品牌", "易被替代", "香港主体 + Baidu 背书 + SLA"],
            ["合规与数据出境", "企业单难签", "先做开发者，企业单单评"],
        ],
        top=Inches(1.55),
        row_h=0.75,
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "资本用途（若融资）", "轻资产：钱用于获客与缓冲，不用于买 GPU")
    table_slide(
        s,
        ["用途", "示意占比", "说明"],
        [
            ["海外获客与渠道", "40%", "开发者社群、内容、渠道分成"],
            ["合规/合同/法务", "15%", "转售授权、合同模板"],
            ["收款与风控", "15%", "Stripe、拒付与盗刷防控"],
            ["客服与技术对接", "15%", "一线支持与上游对接"],
            ["运营现金缓冲", "15%", "覆盖结算周期与波动"],
        ],
        top=Inches(1.7),
        row_h=0.75,
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "投资人应追问的 4 个问题")
    qs = [
        ("01", "百度侧是否允许子账号转售给境外客户？"),
        ("02", "折扣锁定期与调价机制是什么？"),
        ("03", "首批 10 个客户从哪来，谁负责成交？"),
        ("04", "若百度谈崩，自建网关的启动成本与时间？"),
    ]
    for i, (n, q) in enumerate(qs):
        top = Inches(1.5 + i * 1.25)
        card(s, Inches(0.7), top, Inches(11.9), Inches(1.05))
        tb = textbox(s, Inches(1.0), top + Inches(0.28), Inches(1.2), Inches(0.5))
        tf = tb.text_frame
        clear_tf(tf)
        add_text(tf, n, size=20, bold=True, color=COPPER, space_after=0)
        tb2 = textbox(s, Inches(2.4), top + Inches(0.32), Inches(9.8), Inches(0.5))
        tf2 = tb2.text_frame
        clear_tf(tf2)
        add_text(tf2, q, size=16, color=INK2, space_after=0)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    band(s, Inches(0), Inches(7.5), WHITE)
    section_header(s, "结论")
    two_col_cards(
        s,
        "优先路径",
        [
            "• 百度子账号转售（优于纯返点）",
            "• 香港公司收款",
            "• 先开发者/小团队，再企业",
            "• 售价 = 成本 × 1.5～1.8",
        ],
        "明确边界",
        [
            "• 忌：裸卖 Key / 共享账号",
            "• 暂不买服务器",
            "• One API 仅作备用",
            "• 先书面授权，再对外报价",
        ],
        left_color=TEAL,
        right_color=BAD,
    )

    title_slide(
        prs,
        "谢谢",
        "下一动作：把百度五句话谈成书面记录 → 出报价骨架 → 种子客户试单",
        "配套：Cursor 投资者看板 Canvas · 仓库 docs/海外模型与算力业务/",
    )

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
